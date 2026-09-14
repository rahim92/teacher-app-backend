"""On-demand file generation: PDF (single documents) and Excel (bulk export).

Single-document exports (one student's note card, one classroom's grade
sheet) are meant to also be generatable directly on-device in Flutter for
true offline use -- this server-side route exists for parity/bulk cases and
for any client that prefers to fetch a ready-made file.

WeasyPrint and openpyxl are both imported lazily inside their handlers (not
at module import time) so the API can still boot even in an environment
where a native dependency (cairo/pango for WeasyPrint) isn't installed --
only that one export fails until it is. openpyxl itself is already a hard
dependency (see requirements.txt, added for §35's roster import), so its
import basically never fails in practice -- kept lazy anyway for consistency
and because there's no reason to pay the import cost on every app boot for
an endpoint most requests never call.
"""
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Response
from jinja2 import Environment, select_autoescape
from sqlmodel import Session, select

from app.auth import get_current_user
from app.database import get_session
from app.models.assessment import Assessment, AssessmentScore
from app.models.common import AssessmentType, NotebookQuality, SessionEventType
from app.models.identity import Classroom, Subject, Term, User
from app.models.session import ClassSession, SessionEvent
from app.models.student import NotebookCheck, Student
from app.routers.council import _ensure_can_view_council
from app.routers.behavior import compute_behavior_score
from app.routers.students import _ensure_can_manage_student

router = APIRouter(prefix="/export", tags=["export"])

_STUDENT_CARD_TEMPLATE = """
<html dir="rtl" lang="ar">
<head>
<meta charset="utf-8" />
<style>
  body { font-family: "DejaVu Sans", sans-serif; padding: 24px; }
  h1 { font-size: 20px; border-bottom: 2px solid #333; padding-bottom: 8px; }
  .row { margin: 6px 0; }
  .label { font-weight: bold; }
</style>
</head>
<body>
  <h1>بطاقة ملاحظة تلميذ</h1>
  <div class="row"><span class="label">الاسم الكامل:</span> {{ student.first_name }} {{ student.last_name }}</div>
  <div class="row"><span class="label">تاريخ الميلاد:</span> {{ student.birth_date or "-" }}</div>
  <div class="row"><span class="label">ملاحظات:</span> {{ student.general_notes or "-" }}</div>
</body>
</html>
"""


@router.get("/students/{student_id}/card.pdf")
def export_student_card(
    student_id: str,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    student = session.get(Student, student_id)
    if not student or student.is_deleted:
        raise HTTPException(status_code=404, detail="التلميذ غير موجود")
    # Same relation check as editing this student (see students.py):
    # this endpoint takes a bare student_id with no classroom context, and
    # the card includes the student's name/birth date/notes -- unlike the
    # deliberately open cross-teacher GETs elsewhere (ledger, behavior
    # score), there's no classroom_id the caller had to already know here.
    _ensure_can_manage_student(session, student, current_user)

    try:
        from weasyprint import HTML  # lazy import, see module docstring
    except ImportError as exc:  # pragma: no cover - environment-dependent
        raise HTTPException(
            status_code=501,
            detail="توليد PDF على السيرفر يتطلب تثبيت WeasyPrint ومكتباته (cairo/pango).",
        ) from exc

    env = Environment(autoescape=select_autoescape(["html"]))
    html = env.from_string(_STUDENT_CARD_TEMPLATE).render(student=student)
    pdf_bytes = HTML(string=html).write_pdf()

    return Response(content=pdf_bytes, media_type="application/pdf")


# Arabic labels for the raw/bulk export below -- the frontend owns its own
# copies of these for the screens it renders (TAP_TYPES in
# classroom_dashboard.html, etc.), but this file has no frontend to borrow
# labels from, and the raw session-event log below deliberately includes
# every SessionEventType (not just the 8 that have an active tap button
# today, see behavior.py's module docstring) so a row is never a bare
# English enum value even for older/inactive types.
_ASSESSMENT_TYPE_LABELS = {
    AssessmentType.test: "فرض",
    AssessmentType.exam: "اختبار",
    AssessmentType.homework: "عمل منزلي",
    AssessmentType.oral: "شفوي",
}
_NOTEBOOK_QUALITY_LABELS = {
    NotebookQuality.organized: "منظم",
    NotebookQuality.average: "متوسط",
    NotebookQuality.neglected: "مهمل",
}
_SESSION_EVENT_TYPE_LABELS = {
    SessionEventType.attendance_present: "حاضر",
    SessionEventType.attendance_absent: "غائب",
    SessionEventType.tardiness: "تأخر",
    SessionEventType.behavior_positive: "سلوك إيجابي",
    SessionEventType.behavior_negative: "سلوك سلبي",
    SessionEventType.homework_done: "قام بعمل إضافي",
    SessionEventType.homework_missing: "لم يقم بعمل إضافي",
    SessionEventType.participation: "مشاركة / فعالية",
    SessionEventType.equipment_brought: "أحضر الأدوات",
    SessionEventType.equipment_missing: "نسي الأدوات",
    SessionEventType.teamwork_positive: "عمل جماعي جيد",
    SessionEventType.teamwork_negative: "عمل جماعي ضعيف",
    SessionEventType.initiative_shown: "مبادرة",
}


@router.get("/classrooms/{classroom_id}/full-backup.xlsx")
def export_classroom_backup(
    classroom_id: str,
    term_id: str,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """نسخة احتياطية/تصدير شامل (§39): ملف Excel واحد بخمس أوراق (التلاميذ،
    علامة السلوك، الفروض والاختبارات، سجل الحضور والسلوك الخام، فحص الكراس)
    لكل بيانات قسم معيَّن في فصل دراسي معيَّن -- أمان ضد فقدان البيانات،
    وأرشيف نهاية الفصل قابل للفتح خارج التطبيق (Excel/LibreOffice/Google
    Sheets). بنفس بوابة صلاحية تقرير مجلس القسم بالضبط (الأستاذ الرئيسي أو
    الإدارة فقط -- انظر _ensure_can_view_council في routers/council.py)،
    لأن هذا الملف يكشف بيانات كل تلاميذ القسم عبر كل المواد معاً، تماماً
    كتقرير المجلس -- لا يُتاح لأستاذ مادة واحدة فقط.
    """
    classroom = session.get(Classroom, classroom_id)
    if not classroom or classroom.is_deleted:
        raise HTTPException(status_code=404, detail="القسم غير موجود")
    _ensure_can_view_council(classroom, current_user)
    term = session.get(Term, term_id)
    if not term or term.is_deleted:
        raise HTTPException(status_code=404, detail="الفصل الدراسي غير موجود")

    from openpyxl import Workbook  # lazy import, see module docstring
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    students = session.exec(
        select(Student)
        .where(Student.classroom_id == classroom_id, Student.is_deleted == False)  # noqa: E712
        .order_by(Student.first_name, Student.last_name)
    ).all()
    student_name = {s.id: f"{s.first_name} {s.last_name}" for s in students}
    student_ids = list(student_name.keys())
    subject_name = {s.id: s.name for s in session.exec(select(Subject)).all()}

    wb = Workbook()
    header_font = Font(bold=True)

    def write_sheet(ws, headers, rows):
        # RTL sheet view + a bold header row + a readable column width --
        # every sheet in this workbook gets the same treatment, so this is
        # factored out instead of repeated five times.
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
        for row in rows:
            ws.append(row)
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(40, len(str(header)) + 6))

    # -- 1) التلاميذ --
    ws_students = wb.active
    ws_students.title = "التلاميذ"
    write_sheet(
        ws_students,
        ["الاسم", "اللقب", "تاريخ الميلاد", "الجنس", "ولي الأمر", "هاتف ولي الأمر", "ملاحظات طبية", "ملاحظات عامة"],
        [
            [s.first_name, s.last_name, s.birth_date or "", s.gender or "",
             s.guardian_name or "", s.guardian_phone or "", s.medical_notes or "", s.general_notes or ""]
            for s in students
        ],
    )

    # -- 2) علامة السلوك -- نفس compute_behavior_score التي تستعملها شاشة
    # المتابعة المستمرة ومعدل المادة الفصلي، لا حساب موازٍ قد يختلف عنها.
    behavior_scores = [compute_behavior_score(session, s.id, classroom_id, term_id) for s in students]
    category_labels = [b.label_ar for b in behavior_scores[0].breakdown] if behavior_scores else []
    ws_behavior = wb.create_sheet("علامة السلوك")
    write_sheet(
        ws_behavior,
        ["التلميذ"] + category_labels + ["المجموع"],
        [
            [student_name[score.student_id]] + [b.points_earned for b in score.breakdown] + [score.total]
            for score in behavior_scores
        ],
    )

    # -- 3) الفروض والاختبارات -- صف واحد لكل (تلميذ × فرض) هذا الفصل، بما
    # في ذلك التلاميذ الذين لم تُدخَل علامتهم بعد (خانة العلامة تبقى فارغة
    # بدل حذف الصف بصمت -- هذا بالضبط ما تُبلِّغ عنه تنبيهات §38).
    assessments = session.exec(
        select(Assessment).where(
            Assessment.classroom_id == classroom_id,
            Assessment.date >= term.start_date,
            Assessment.date <= term.end_date,
            Assessment.is_deleted == False,  # noqa: E712
        ).order_by(Assessment.date)
    ).all()
    assessment_rows = []
    for assessment in assessments:
        scores = session.exec(
            select(AssessmentScore).where(
                AssessmentScore.assessment_id == assessment.id,
                AssessmentScore.is_deleted == False,  # noqa: E712
            )
        ).all()
        score_by_student = {sc.student_id: sc.score for sc in scores}
        for s in students:
            assessment_rows.append([
                student_name[s.id],
                subject_name.get(assessment.subject_id, "—"),
                assessment.title,
                _ASSESSMENT_TYPE_LABELS.get(assessment.assessment_type, assessment.assessment_type),
                assessment.date,
                score_by_student.get(s.id),
                assessment.max_score,
            ])
    ws_assessments = wb.create_sheet("الفروض والاختبارات")
    write_sheet(ws_assessments, ["التلميذ", "المادة", "العنوان", "النوع", "التاريخ", "العلامة", "من"], assessment_rows)

    # -- 4) سجل الحضور والسلوك -- كل نقرة خام هذا الفصل، للأرشفة/التدقيق.
    events_with_sessions = (
        session.exec(
            select(SessionEvent, ClassSession)
            .join(ClassSession, SessionEvent.session_id == ClassSession.id)  # type: ignore[arg-type]
            .where(
                ClassSession.classroom_id == classroom_id,
                ClassSession.date >= term.start_date,
                ClassSession.date <= term.end_date,
                SessionEvent.is_deleted == False,  # noqa: E712
            )
            .order_by(ClassSession.date)
        ).all()
        if student_ids
        else []
    )
    log_rows = [
        [
            class_session.date,
            subject_name.get(class_session.subject_id, "—"),
            student_name.get(event.student_id, "؟"),
            _SESSION_EVENT_TYPE_LABELS.get(event.event_type, event.event_type),
            event.note or "",
        ]
        for event, class_session in events_with_sessions
        if event.student_id in student_name
    ]
    ws_log = wb.create_sheet("سجل الحضور والسلوك")
    write_sheet(ws_log, ["التاريخ", "المادة", "التلميذ", "النوع", "ملاحظة"], log_rows)

    # -- 5) فحص الكراس --
    checks = (
        session.exec(
            select(NotebookCheck).where(
                NotebookCheck.student_id.in_(student_ids),  # type: ignore[union-attr]
                NotebookCheck.check_date >= term.start_date,
                NotebookCheck.check_date <= term.end_date,
                NotebookCheck.is_deleted == False,  # noqa: E712
            ).order_by(NotebookCheck.check_date)
        ).all()
        if student_ids
        else []
    )
    ws_notebook = wb.create_sheet("فحص الكراس")
    write_sheet(
        ws_notebook,
        ["التلميذ", "التاريخ", "تنظيم الكراس", "الكتابة", "ملاحظة"],
        [
            [
                student_name.get(check.student_id, "؟"),
                check.check_date,
                _NOTEBOOK_QUALITY_LABELS.get(check.quality, check.quality),
                _NOTEBOOK_QUALITY_LABELS.get(check.writing_quality, "") if check.writing_quality else "",
                check.note or "",
            ]
            for check in checks
        ],
    )

    buffer = BytesIO()
    wb.save(buffer)
    # Deliberately ASCII-only (no Arabic classroom/term name): a
    # Content-Disposition filename has to be Latin-1-safe or Starlette
    # raises encoding a UnicodeEncodeError building the response headers.
    # The frontend gives the download its real Arabic name itself via the
    # <a download="..."> attribute on the blob URL it builds (see
    # downloadClassroomBackup() in classroom_dashboard.html) -- that name,
    # not this one, is what the teacher actually sees saved to disk.
    filename = f"classroom-backup-{classroom_id}-{term_id}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
